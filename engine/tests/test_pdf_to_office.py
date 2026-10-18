"""PDF → Excel (D4), PDF → PowerPoint (D3) and PDF → Text (D6)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pikepdf
import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas

from jobs import Cancelled
from protocol import EngineError, ErrorCode
from tests.fixtures import make_corrupt_pdf, make_document_pdf, make_encrypted_pdf
from tests.helpers import RecordingContext, sha256
from tools.pdf_office import (
    _strip_text_operators,
    cell_value,
    pdf_to_excel,
    pdf_to_powerpoint,
    pdf_to_text,
)


def run(tool: Any, ctx: RecordingContext, inputs: list[Path], **options: Any) -> Any:
    return tool({"inputs": [str(p) for p in inputs], "options": options}, ctx)


def make_columns_pdf(path: Path) -> Path:
    """An unruled, column-aligned price list (no table lines at all)."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(72, 760, "Price list")
    rows = [("Item", "Qty", "Price"), ("Apples", "12", "1,250.50"), ("Pears", "7", "99")]
    for r, (item, qty, price) in enumerate(rows):
        y = 720 - r * 18
        c.setFont("Helvetica-Bold" if r == 0 else "Helvetica", 11)
        c.drawString(72, y, item)
        c.drawString(250, y, qty)
        c.drawString(350, y, price)
    c.showPage()
    c.save()
    return path


# ── Excel ─────────────────────────────────────────────────────────────────────────────────


def test_cell_values() -> None:
    assert cell_value("120") == (120, None)
    assert cell_value("1,250.50") == (1250.5, "#,##0.##")
    assert cell_value("-3.5") == (-3.5, None)
    assert cell_value("12%") == (0.12, "0.##%")
    assert cell_value("12-05-2026") == ("12-05-2026", None)
    assert cell_value("Apples") == ("Apples", None)


def test_ruled_tables_one_sheet_each(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=2)
    before = sha256(pdf)
    result = run(pdf_to_excel, ctx, [pdf])
    out = result.outputs[0]
    assert out == tmp_path / "report.xlsx"
    assert sha256(pdf) == before
    wb = load_workbook(out)
    assert wb.sheetnames == ["Page 1 Table 1"]
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    assert rows == [("Region", "Q1", "Q2"), ("North", 120, 135), ("South", 98, 101)]
    assert wb.worksheets[0]["A1"].font.bold
    assert result.stats["tables"] == 1


def test_without_ruled_tables_text_goes_into_columns(tmp_path: Path, ctx: RecordingContext) -> None:
    result = run(pdf_to_excel, ctx, [make_columns_pdf(tmp_path / "prices.pdf")])
    wb = load_workbook(result.outputs[0])
    assert wb.sheetnames == ["Page 1"]  # fell back to the per-page layout
    rows = [r for r in wb.worksheets[0].iter_rows(values_only=True) if any(r)]
    assert rows[0][0] == "Price list"
    assert rows[1] == ("Item", "Qty", "Price")
    assert rows[2] == ("Apples", 12, 1250.5)
    assert rows[3] == ("Pears", 7, 99)


def test_one_sheet_per_page_and_single_sheet(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=3)
    pages = load_workbook(run(pdf_to_excel, ctx, [pdf], layout="pages").outputs[0])
    assert pages.sheetnames == ["Page 1", "Page 2", "Page 3"]
    single_result = run(pdf_to_excel, RecordingContext(tmp_path / "j2"), [pdf], layout="single")
    single = load_workbook(single_result.outputs[0])
    assert single.sheetnames == ["report"]
    values = [v for row in single.worksheets[0].iter_rows(values_only=True) for v in row if v]
    assert values.count("Region") == 1
    assert sum(1 for v in values if str(v).startswith("Chapter")) == 3


def test_excel_encrypted_and_damaged(tmp_path: Path, ctx: RecordingContext) -> None:
    secret = make_encrypted_pdf(tmp_path / "secret.pdf", 1, "pw")
    with pytest.raises(EngineError) as err:
        run(pdf_to_excel, ctx, [secret])
    assert err.value.code == ErrorCode.PASSWORD_REQUIRED
    with pytest.raises(EngineError) as err:
        broken = make_corrupt_pdf(tmp_path / "broken.pdf")
        run(pdf_to_excel, RecordingContext(tmp_path / "j2"), [broken])
    assert err.value.code == ErrorCode.CORRUPT_FILE


# ── PowerPoint ────────────────────────────────────────────────────────────────────────────


def test_one_slide_per_page_with_text_boxes(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=2)
    result = run(pdf_to_powerpoint, ctx, [pdf])
    out = result.outputs[0]
    assert out == tmp_path / "report.pptx"
    deck = Presentation(str(out))
    assert len(deck.slides) == 2
    assert abs(deck.slide_width - 595.27 * 12700) < 12700  # A4 portrait slide
    first = list(deck.slides[0].shapes)
    assert first[0].shape_type == MSO_SHAPE_TYPE.PICTURE  # background
    texts = [s.text_frame.text for s in first if s.has_text_frame]
    assert "Annual Report 2026" in texts
    assert any("bold words, italic words" in t for t in texts)
    bold = [
        r
        for s in first
        if s.has_text_frame
        for p in s.text_frame.paragraphs
        for r in p.runs
        if r.text.strip() == "bold words"
    ]
    assert bold and bold[0].font.bold


def test_background_has_no_text_but_keeps_drawings(tmp_path: Path) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=1)
    with pikepdf.open(pdf) as doc:
        _strip_text_operators(doc)
        ops = [str(i.operator) for i in pikepdf.parse_content_stream(doc.pages[0])]
    assert "BT" not in ops and "Tj" not in ops and "TJ" not in ops
    assert "re" in ops or "l" in ops  # table rules survive
    assert "Do" in ops  # the picture survives


def test_mixed_page_sizes_fit_into_the_deck(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = tmp_path / "mixed.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    c.drawString(72, 700, "Portrait page")
    c.showPage()
    c.setPageSize(landscape(A4))
    c.drawString(72, 500, "Landscape page")
    c.showPage()
    c.save()
    deck = Presentation(str(run(pdf_to_powerpoint, ctx, [pdf]).outputs[0]))
    assert len(deck.slides) == 2
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + 12700


def test_cancel_powerpoint(tmp_path: Path) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=3)
    with pytest.raises(Cancelled):
        run(pdf_to_powerpoint, RecordingContext(tmp_path / "job", cancel_after=2), [pdf])
    assert not (tmp_path / "report.pptx").exists()


# ── Text ──────────────────────────────────────────────────────────────────────────────────


def test_text_export(tmp_path: Path, ctx: RecordingContext) -> None:
    result = run(pdf_to_text, ctx, [make_document_pdf(tmp_path / "report.pdf", pages=2)])
    out = result.outputs[0]
    assert out == tmp_path / "report.txt"
    raw = out.read_bytes()
    assert b"\r\n" in raw
    text = raw.decode("utf-8")
    pages = text.split("\f")
    assert len(pages) == 2
    assert pages[0].startswith("Annual Report 2026")
    assert "bold words, italic words and red words." in text
    assert result.stats["pages"] == 2


def test_text_keep_layout_preserves_columns(tmp_path: Path, ctx: RecordingContext) -> None:
    out = run(pdf_to_text, ctx, [make_columns_pdf(tmp_path / "prices.pdf")], keep_layout=True)
    lines = out.outputs[0].read_text(encoding="utf-8").splitlines()
    apples = next(line for line in lines if "Apples" in line)
    header = next(line for line in lines if "Item" in line)
    assert apples.index("12") == header.index("Qty")  # columns line up
