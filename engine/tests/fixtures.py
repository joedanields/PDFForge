"""Generators for test inputs (docs/07_TESTING.md §2). Everything is created on the fly so no
binary corpus is committed; later phases extend this module.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any, Literal

import pikepdf
from pikepdf import Array, Dictionary, Name, NameTree, OutlineItem
from PIL import Image
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas


def page_label(label: str, n: int) -> str:
    """Text drawn on page ``n`` (1-based) of a fixture called ``label``."""
    return f"{label} page {n}"


def make_text_pdf(
    path: Path,
    pages: int,
    label: str | None = None,
    *,
    bookmarks: bool = False,
    named_link: bool = False,
    page_sizes: Sequence[tuple[float, float]] | None = None,
) -> Path:
    """A simple text PDF; each page shows ``"{label} page {n}"``.

    ``bookmarks``: one outline entry per page ("{label} §n") using named destinations.
    ``named_link``: page 1 gets an internal link (named destination) to the last page.
    """
    label = label or path.stem
    c = canvas.Canvas(str(path), pagesize=A4)
    for n in range(1, pages + 1):
        size = page_sizes[(n - 1) % len(page_sizes)] if page_sizes else A4
        c.setPageSize(size)
        c.setFont("Helvetica", 24)
        c.drawString(72, size[1] - 100, page_label(label, n))
        key = f"p{n}"
        c.bookmarkPage(key)
        if bookmarks:
            c.addOutlineEntry(f"{label} §{n}", key, level=0)
        if named_link and n == 1 and pages > 1:
            c.linkRect("", f"p{pages}", (72, 400, 300, 440), relative=0)
        c.showPage()
    c.save()
    return path


LinkStyle = Literal["explicit", "named", "goto", "legacy"]


def _name_tree(pdf: pikepdf.Pdf) -> NameTree:
    names = pdf.Root.get("/Names")
    if not isinstance(names, Dictionary):
        names = pdf.make_indirect(Dictionary())
        pdf.Root.Names = names
    if "/Dests" not in names:
        names.Dests = NameTree.new(pdf).obj
    return NameTree(names.Dests)


def _named_target(pdf: pikepdf.Pdf, name: str, to_page: int, legacy: bool) -> Any:
    explicit = Array([pdf.pages[to_page].obj, Name.Fit])
    if legacy:
        dests = pdf.Root.get("/Dests")
        if not isinstance(dests, Dictionary):
            dests = pdf.make_indirect(Dictionary())
            pdf.Root.Dests = dests
        dests[Name(f"/{name}")] = explicit
        return Name(f"/{name}")
    _name_tree(pdf)[name] = Dictionary(D=explicit)
    return pikepdf.String(name)


def add_link(path: Path, from_page: int, to_page: int, style: LinkStyle = "explicit") -> Path:
    """Add an internal Link annotation (0-based pages).

    ``explicit``: ``/Dest [page /Fit]``; ``named``: ``/Dest (name)`` via the /Names tree;
    ``goto``: ``/A << /S /GoTo /D (name) >>``; ``legacy``: ``/Dest /name`` via PDF 1.1 /Dests.
    """
    with pikepdf.open(path, allow_overwriting_input=True) as pdf:
        fields: dict[str, Any]
        if style == "explicit":
            fields = {"Dest": Array([pdf.pages[to_page].obj, Name.Fit])}
        else:
            target = _named_target(pdf, f"d{from_page}to{to_page}", to_page, style == "legacy")
            fields = (
                {"A": Dictionary(S=Name.GoTo, D=target)} if style == "goto" else {"Dest": target}
            )
        annot = pdf.make_indirect(
            Dictionary(
                Type=Name.Annot,
                Subtype=Name.Link,
                Rect=Array([72, 300, 300, 340]),
                Border=Array([0, 0, 0]),
                **fields,
            )
        )
        page = pdf.pages[from_page].obj
        annots = page.get("/Annots")
        if isinstance(annots, Array):
            annots.append(annot)
        else:
            page.Annots = Array([annot])
        pdf.save(path)
    return path


def add_action_outline(
    path: Path, entries: Sequence[tuple[str, int]], uri: str | None = None
) -> Path:
    """Replace the outline with GoTo-action bookmarks to named destinations, plus an optional
    URI bookmark (used to check non-page actions survive)."""
    with pikepdf.open(path, allow_overwriting_input=True) as pdf:
        actions = [
            (title, Dictionary(S=Name.GoTo, D=_named_target(pdf, f"o{n}", page, legacy=False)))
            for n, (title, page) in enumerate(entries)
        ]
        with pdf.open_outline() as outline:
            outline.root.clear()
            for title, action in actions:
                outline.root.append(OutlineItem(title, action=action))
            if uri:
                outline.root.append(OutlineItem("Website", action=Dictionary(S=Name.URI, URI=uri)))
        pdf.save(path)
    return path


def make_landscape_pdf(path: Path, pages: int, label: str | None = None) -> Path:
    return make_text_pdf(path, pages, label, page_sizes=[landscape(A4)])


def make_encrypted_pdf(
    path: Path, pages: int, user_password: str, owner_password: str = ""
) -> Path:
    plain = path.with_name(f"{path.stem}.plain.pdf")
    make_text_pdf(plain, pages, path.stem)
    with pikepdf.open(plain) as pdf:
        pdf.save(
            path,
            encryption=pikepdf.Encryption(
                user=user_password, owner=owner_password or user_password, R=6
            ),
        )
    plain.unlink()
    return path


def make_corrupt_pdf(path: Path) -> Path:
    """Starts like a PDF, contains no recoverable structure."""
    path.write_bytes(b"%PDF-1.7\n" + b"\x00garbage\xff" * 200)
    return path


def make_not_pdf(path: Path) -> Path:
    path.write_text("this is plain text, not a PDF\n", encoding="utf-8")
    return path


def make_image(
    path: Path,
    size: tuple[int, int] = (400, 300),
    mode: str = "RGB",
    color: tuple[int, ...] = (30, 120, 200),
) -> Path:
    """An image in the format implied by ``path``'s suffix."""
    img = Image.new(mode, size, color if len(color) == len(mode) else color[: len(mode)])
    img.save(path)
    return path


def make_form_pdf(path: Path, field_name: str = "name") -> Path:
    """One page with a single AcroForm text field."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 760, f"Form {path.stem}")
    c.acroForm.textfield(name=field_name, x=72, y=700, width=200, height=24)
    c.showPage()
    c.save()
    return path


def make_restricted_pdf(path: Path, pages: int, owner_password: str = "owner-secret") -> Path:
    """Opens without a password but carries owner restrictions (no printing, copying, editing)."""
    plain = path.with_name(f"{path.stem}.plain.pdf")
    make_text_pdf(plain, pages, path.stem)
    restrictions = pikepdf.Permissions(
        extract=False,
        modify_other=False,
        modify_annotation=False,
        print_lowres=False,
        print_highres=False,
        modify_assembly=False,
        modify_form=False,
    )
    with pikepdf.open(plain) as pdf:
        pdf.save(
            path, encryption=pikepdf.Encryption(user="", owner=owner_password, allow=restrictions)
        )
    plain.unlink()
    return path


def make_image_pdf(
    path: Path,
    pages: int,
    pixels: tuple[int, int] = (800, 1000),
    *,
    seed: int = 1,
    quality: int = 92,
    grain: int = 4,
) -> Path:
    """A 'scanned' PDF: each A4 page is one noisy RGB photo (hard to compress), JPEG-encoded."""
    import io
    import random

    from reportlab.lib.utils import ImageReader

    rng = random.Random(seed)
    c = canvas.Canvas(str(path), pagesize=A4)
    for n in range(1, pages + 1):
        small = (max(1, pixels[0] // grain), max(1, pixels[1] // grain))
        noise = rng.randbytes(small[0] * small[1] * 3)
        img = Image.frombytes("RGB", small, noise).resize(pixels, Image.Resampling.BILINEAR)
        buffer = io.BytesIO()
        img.save(buffer, "JPEG", quality=quality)
        buffer.seek(0)
        c.drawImage(ImageReader(buffer), 36, 72, width=A4[0] - 72, height=A4[1] - 144)
        c.setFont("Helvetica", 14)
        c.drawString(36, 40, page_label(path.stem, n))
        c.showPage()
    c.save()
    return path


def make_exif_rotated_jpeg(
    path: Path, size: tuple[int, int] = (400, 200), orientation: int = 6
) -> Path:
    """Stored landscape, displayed portrait (EXIF 6 = rotate 90° clockwise to display)."""
    img = Image.new("RGB", size, (200, 40, 40))
    exif = Image.Exif()
    exif[0x0112] = orientation
    img.save(path, "JPEG", exif=exif)
    return path


def make_multipage_tiff(path: Path, frames: int = 3, size: tuple[int, int] = (300, 400)) -> Path:
    pages = [Image.new("RGB", size, (40 * n, 80, 120)) for n in range(frames)]
    pages[0].save(path, "TIFF", save_all=True, append_images=pages[1:])
    return path


def make_heic(path: Path, size: tuple[int, int] = (320, 240)) -> Path:
    """HEIC test image, encoded with pillow-heif (dev-only; the app only decodes, via pi-heif)."""
    import pillow_heif

    heif = pillow_heif.from_pillow(Image.new("RGB", size, (20, 160, 90)))
    heif.save(str(path), quality=80)
    return path


# ── Office documents (Phase 2 conversions) ────────────────────────────────────────────────


def make_docx(path: Path, pages: int = 3, title: str = "Quarterly report") -> Path:
    """DOCX with a heading, ``pages`` pages of paragraphs (explicit page breaks) and a table."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    for n in range(1, pages + 1):
        doc.add_paragraph(f"Section {n}. " + "The quick brown fox jumps over the lazy dog. " * 6)
        if n == 1:
            table = doc.add_table(rows=2, cols=2)
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    cell.text = f"R{r + 1}C{c + 1}"
        if n < pages:
            doc.add_page_break()  # type: ignore[no-untyped-call]
    doc.save(str(path))
    return path


def make_pptx(path: Path, slides: int = 3, notes: bool = True) -> Path:
    from pptx import Presentation

    deck = Presentation()
    for n in range(1, slides + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Slide {n}"
        if notes:
            slide.notes_slide.notes_text_frame.text = f"Speaker note {n}"
    deck.save(str(path))
    return path


def make_xlsx(path: Path, columns: int = 30, rows: int = 20, active: int = 0) -> Path:
    """Workbook with a wide 'Data' sheet and a small 'Summary' sheet; ``active`` = active tab."""
    from openpyxl import Workbook

    wb = Workbook()
    data = wb.worksheets[0]
    data.title = "Data"
    data.append([f"Column {c}" for c in range(1, columns + 1)])
    for r in range(1, rows + 1):
        data.append([r * c for c in range(1, columns + 1)])
    summary = wb.create_sheet("Summary")
    summary.append(["Summary", "", "total"])
    wb.active = active
    wb.save(str(path))
    return path


def make_encrypted_office(path: Path) -> Path:
    """What a password-protected DOCX/XLSX/PPTX looks like: an OLE compound file."""
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 504)
    return path


def make_document_pdf(path: Path, pages: int = 2) -> Path:
    """A typical digitally created report: title, headings, wrapped paragraphs with bold /
    italic / coloured words, a ruled table and a picture (for PDF → Word / Excel / PPT)."""
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Image as RLImage,
    )
    from reportlab.platypus import (
        PageBreak,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.platypus import (
        Paragraph as RLParagraph,
    )

    styles = getSampleStyleSheet()
    picture = path.with_suffix(".picture.png")
    Image.new("RGB", (240, 160), (30, 120, 200)).save(picture)
    story: list[Any] = [RLParagraph("Annual Report 2026", styles["Title"])]
    for n in range(1, pages + 1):
        story.append(RLParagraph(f"Chapter {n} overview", styles["Heading1"]))
        story.append(
            RLParagraph(
                "This paragraph contains <b>bold words</b>, <i>italic words</i> and "
                '<font color="#C00000">red words</font>. '
                + "The quick brown fox jumps over the lazy dog. " * 8,
                styles["BodyText"],
            )
        )
        story.append(
            RLParagraph("Second paragraph. " + "Lorem ipsum dolor sit. " * 10, styles["BodyText"])
        )
        if n == 1:
            table = Table(
                [["Region", "Q1", "Q2"], ["North", "120", "135"], ["South", "98", "101"]],
                colWidths=[120, 60, 60],
            )
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
            story += [Spacer(1, 12), table, Spacer(1, 12), RLImage(str(picture), 120, 80)]
        if n < pages:
            story.append(PageBreak())
    SimpleDocTemplate(str(path)).build(story)
    picture.unlink()
    return path


def make_scanned_text_pdf(
    path: Path,
    lines: Sequence[str] = ("Invoice number 4711", "Total amount due 1234"),
    pages: int = 1,
    dpi: int = 200,
    skew: float = 0.0,
    turn: int = 0,
) -> Path:
    """An image-only 'scan' (A4 at ``dpi``) of large black text: no text layer at all."""
    import img2pdf
    from PIL import ImageDraw, ImageFont

    size = (round(8.27 * dpi), round(11.69 * dpi))
    try:
        font: Any = ImageFont.truetype("arial.ttf", round(dpi * 0.28))
    except OSError:
        font = ImageFont.load_default(size=round(dpi * 0.28))
    frames = []
    for n in range(pages):
        img = Image.new("L", size, 255)
        draw = ImageDraw.Draw(img)
        for i, text in enumerate(lines):
            draw.text(
                (dpi, dpi + i * dpi * 0.6),
                f"{text}" if n == 0 else f"{text} p{n + 1}",
                fill=0,
                font=font,
            )
        if skew:
            img = img.rotate(skew, resample=Image.Resampling.BICUBIC, fillcolor=255, expand=False)
        if turn:  # scanned sideways: the page shows the text rotated by ``turn`` degrees CCW
            img = img.rotate(turn, expand=True)
        buffer = path.with_suffix(f".{n}.png")
        img.save(buffer, dpi=(dpi, dpi))
        frames.append(buffer)
    path.write_bytes(img2pdf.convert([str(f) for f in frames]))
    for f in frames:
        f.unlink()
    return path


def make_full_form_pdf(path: Path) -> Path:
    """Page 1: name (text), notes (multi-line text), agree (check box), size (radio
    small/large), country (combo box). Page 2: city (text)."""
    c = canvas.Canvas(str(path), pagesize=A4)
    form = c.acroForm
    c.drawString(72, 790, "Application form")
    form.textfield(name="name", x=72, y=740, width=250, height=22)
    form.textfield(name="notes", x=72, y=640, width=300, height=80, fieldFlags="multiline")
    form.checkbox(name="agree", x=72, y=600, size=16, buttonStyle="check")
    form.radio(name="size", value="small", x=72, y=560, size=16, selected=True)
    form.radio(name="size", value="large", x=120, y=560, size=16)
    form.choice(
        name="country",
        x=72,
        y=500,
        width=200,
        height=22,
        options=["India", "Japan", "Peru"],
        value="India",
    )
    c.showPage()
    c.acroForm.textfield(name="city", x=72, y=740, width=250, height=22)
    c.showPage()
    c.save()
    return path


def make_pkcs12(
    path: Path, password: str, *, name: str = "Asha Raman", days: int = 365, expired: bool = False
) -> Path:
    """A self-signed signing certificate with its private key, as a PFX/P12 file."""
    import datetime as dt

    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    subject = x509.Name([x509.NameAttribute(x509.NameOID.COMMON_NAME, name)])
    now = dt.datetime.now(dt.UTC)
    start = now - dt.timedelta(days=days + 10) if expired else now - dt.timedelta(days=1)
    end = now - dt.timedelta(days=1) if expired else now + dt.timedelta(days=days)
    cert = (
        x509.CertificateBuilder()
        .subject_name(subject)
        .issuer_name(subject)
        .public_key(key.public_key())
        .serial_number(x509.random_serial_number())
        .not_valid_before(start)
        .not_valid_after(end)
        .add_extension(
            x509.KeyUsage(
                digital_signature=True,
                content_commitment=True,
                key_encipherment=False,
                data_encipherment=False,
                key_agreement=False,
                key_cert_sign=False,
                crl_sign=False,
                encipher_only=False,
                decipher_only=False,
            ),
            critical=True,
        )
        .sign(key, hashes.SHA256())
    )
    path.write_bytes(
        pkcs12.serialize_key_and_certificates(
            name.encode(),
            key,
            cert,
            None,
            serialization.BestAvailableEncryption(password.encode()),
        )
    )
    return path
