"""PDF → Word (docs/02 D2): own converter on pdfplumber + python-docx (AGPL-free)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.shared import Pt

from jobs import Cancelled
from protocol import EngineError, ErrorCode
from tests.fixtures import (
    make_corrupt_pdf,
    make_document_pdf,
    make_encrypted_pdf,
    make_scanned_text_pdf,
)
from tests.helpers import RecordingContext, sha256
from tools.pdf_office import pdf_to_word
from utils.vendor import tesseract_path

needs_tesseract = pytest.mark.skipif(
    not tesseract_path().is_file(), reason="Tesseract not fetched (scripts/fetch-vendor.ps1)"
)


def run(ctx: RecordingContext, inputs: list[Path], **options: Any) -> Any:
    return pdf_to_word({"inputs": [str(p) for p in inputs], "options": options}, ctx)


def runs_by_text(doc: Any) -> dict[str, Any]:
    return {r.text.strip(): r for p in doc.paragraphs for r in p.runs if r.text.strip()}


def test_text_headings_and_formatting(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=2)
    before = sha256(pdf)
    result = run(ctx, [pdf])
    out = result.outputs[0]
    assert out == tmp_path / "report.docx"
    assert sha256(pdf) == before
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts[0] == "Annual Report 2026"
    assert "Chapter 1 overview" in texts
    body = next(t for t in texts if t.startswith("This paragraph contains"))
    # Wrapped lines are joined into one paragraph, punctuation stays attached.
    assert "bold words, italic words and red words." in body
    assert body.count("The quick brown fox") == 8
    headings = [
        p.text for p in doc.paragraphs if p.style and str(p.style.name).startswith("Heading")
    ]
    assert headings[:2] == ["Annual Report 2026", "Chapter 1 overview"]
    runs = runs_by_text(doc)
    assert runs["bold words"].bold
    assert runs["italic words"].italic
    assert str(runs["red words"].font.color.rgb) == "C00000"
    assert result.stats["pages"] == 2


def test_page_size_and_one_section_per_page(tmp_path: Path, ctx: RecordingContext) -> None:
    out = run(ctx, [make_document_pdf(tmp_path / "report.pdf", pages=3)]).outputs[0]
    doc = Document(str(out))
    assert len(doc.sections) == 3
    for section in doc.sections:
        assert section.page_width is not None and section.page_height is not None
        assert abs(section.page_width - Pt(595.27)) < Pt(1)
        assert abs(section.page_height - Pt(841.89)) < Pt(1)


def test_ruled_table_and_picture(tmp_path: Path, ctx: RecordingContext) -> None:
    result = run(ctx, [make_document_pdf(tmp_path / "report.pdf", pages=1)])
    doc = Document(str(result.outputs[0]))
    assert len(doc.tables) == 1
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows == [["Region", "Q1", "Q2"], ["North", "120", "135"], ["South", "98", "101"]]
    # Table text is not repeated as loose paragraphs.
    assert not any(p.text.strip() == "North" for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1
    picture = doc.inline_shapes[0]
    assert abs(picture.width - Pt(120)) < Pt(2)
    assert result.stats["tables"] == 1
    assert result.stats["images"] == 1


def test_scan_without_ocr_keeps_the_page_image(tmp_path: Path, ctx: RecordingContext) -> None:
    result = run(ctx, [make_scanned_text_pdf(tmp_path / "scan.pdf")])
    doc = Document(str(result.outputs[0]))
    assert len(doc.inline_shapes) == 1
    assert not [p.text for p in doc.paragraphs if p.text.strip()]
    assert result.stats["ocr_pages"] == 0


@needs_tesseract
def test_scan_with_ocr_becomes_editable_text(tmp_path: Path, ctx: RecordingContext) -> None:
    result = run(ctx, [make_scanned_text_pdf(tmp_path / "scan.pdf", pages=2)], ocr=True)
    doc = Document(str(result.outputs[0]))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "Invoice number 4711" in texts
    assert "Total amount due 1234 p2" in texts
    assert len(doc.inline_shapes) == 0
    assert result.stats["ocr_pages"] == 2


@needs_tesseract
def test_unknown_ocr_language_is_refused(tmp_path: Path, ctx: RecordingContext) -> None:
    with pytest.raises(EngineError) as err:
        run(ctx, [make_scanned_text_pdf(tmp_path / "scan.pdf")], ocr=True, languages=["xyz"])
    assert err.value.code == ErrorCode.UNSUPPORTED


def test_encrypted_input_needs_the_password(tmp_path: Path, ctx: RecordingContext) -> None:
    pdf = make_encrypted_pdf(tmp_path / "secret.pdf", 2, "pw123")
    with pytest.raises(EngineError) as err:
        run(ctx, [pdf])
    assert err.value.code == ErrorCode.PASSWORD_REQUIRED
    params = {"inputs": [str(pdf)], "passwords": {str(pdf): "pw123"}}
    out = pdf_to_word(params, RecordingContext(tmp_path / "job2")).outputs[0]
    assert out.name == "secret.docx"
    assert any(p.text for p in Document(str(out)).paragraphs)


def test_damaged_pdf(tmp_path: Path, ctx: RecordingContext) -> None:
    with pytest.raises(EngineError) as err:
        run(ctx, [make_corrupt_pdf(tmp_path / "broken.pdf")])
    assert err.value.code == ErrorCode.CORRUPT_FILE


def test_never_overwrites_and_several_files(tmp_path: Path, ctx: RecordingContext) -> None:
    first = make_document_pdf(tmp_path / "a.pdf", pages=1)
    second = make_document_pdf(tmp_path / "b.pdf", pages=1)
    (tmp_path / "a.docx").write_bytes(b"existing")
    result = run(ctx, [first, second])
    assert [p.name for p in result.outputs] == ["a (1).docx", "b.docx"]
    assert (tmp_path / "a.docx").read_bytes() == b"existing"
    assert result.stats["files"] == 2


def test_cancel_leaves_no_output(tmp_path: Path) -> None:
    pdf = make_document_pdf(tmp_path / "report.pdf", pages=3)
    ctx = RecordingContext(tmp_path / "job", cancel_after=2)
    with pytest.raises(Cancelled):
        run(ctx, [pdf])
    assert not (tmp_path / "report.docx").exists()
