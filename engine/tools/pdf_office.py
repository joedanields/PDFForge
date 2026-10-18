"""PDF → Word / Excel / PowerPoint / Text (docs/02_FEATURES.md D2-D4, D6).

AGPL-free by decision (09 §3): our own reconstruction on pdfplumber (layout), pypdfium2
(rendering) and python-docx / openpyxl / python-pptx (writing) — no pdf2docx / PyMuPDF.
Layout fidelity is best effort and the UI says so.
"""

from __future__ import annotations

import io
import logging
import re
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pdfplumber
import pikepdf
import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu
from pptx.util import Pt as PptPt

from jobs import JobContext, ToolResult
from protocol import EngineError, ErrorCode
from utils import tesseract
from utils.layout import (
    Box,
    Paragraph,
    Picture,
    TableBlock,
    Word,
    body_size,
    cluster_columns,
    group_lines,
    group_paragraphs,
    inside,
    split_segments,
    text_frame,
    words_from_plumber,
)
from utils.pdfio import deliver_converted, open_pdf
from utils.validation import (
    check_pdf,
    option_bool,
    option_choice,
    options_dict,
    passwords_map,
    require_str_list,
)

log = logging.getLogger("pdfforge.pdf_office")

CONVERT_MAX = 100
IMAGE_DPI = 150
OCR_DPI = 300
MIN_PICTURE_PT = 12.0
MIN_MARGIN_PT = 18.0
_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# ── shared input handling ─────────────────────────────────────────────────────────────────


def pdf_inputs(params: dict[str, Any]) -> list[Path]:
    paths = [Path(p) for p in require_str_list(params, "inputs", min_len=1, max_len=CONVERT_MAX)]
    for path in paths:
        check_pdf(path)
    return paths


@contextmanager
def readable_copy(source: Path, params: dict[str, Any], work: Path) -> Iterator[tuple[Path, int]]:
    """A path pdfplumber/pdfium can read without a password, and the page count.

    Password and damage problems surface as protocol errors (via pikepdf); encrypted files are
    decrypted into the job folder first.
    """
    with open_pdf(source, passwords_map(params)) as pdf:
        pages = len(pdf.pages)
        if not pdf.is_encrypted:
            yield source, pages
            return
        work.mkdir(parents=True, exist_ok=True)
        plain = work / "decrypted.pdf"
        pdf.save(plain)
    yield plain, pages


def _open_plumber(path: Path, shown: Path) -> Any:
    try:
        return pdfplumber.open(str(path))
    except Exception as exc:  # pdfminer raises many types for broken files
        raise EngineError(
            ErrorCode.CORRUPT_FILE, f"{shown.name} is damaged and could not be read", str(shown)
        ) from exc


def page_words(page: Any) -> list[Word]:
    raw = page.extract_words(
        extra_attrs=["fontname", "size", "non_stroking_color"],
        keep_blank_chars=False,
        use_text_flow=False,
    )
    return words_from_plumber(w for w in raw if w.get("upright", True))


def render_page(doc: pdfium.PdfDocument, index: int, dpi: int) -> Image.Image:
    page = doc[index]
    try:
        bitmap = page.render(scale=dpi / 72, may_draw_forms=True)
        image: Image.Image = bitmap.to_pil().convert("RGB")
        return image
    finally:
        page.close()


def ocr_page_words(
    image: Image.Image, work: Path, langs: list[str], ctx: JobContext, name: str
) -> list[Word]:
    """OCR a rendered page; boxes converted back to points."""
    work.mkdir(parents=True, exist_ok=True)
    png = work / f"{name}.png"
    image.save(png, dpi=(OCR_DPI, OCR_DPI))
    factor = 72 / OCR_DPI
    words = []
    for w in tesseract.ocr_words(png, work / name, langs, OCR_DPI, ctx):
        if w.conf < 30:
            continue
        height = w.height * factor
        words.append(
            Word(
                w.text,
                w.left * factor,
                w.top * factor,
                (w.left + w.width) * factor,
                w.top * factor + height,
                max(6.0, round(height / 0.72 * 2) / 2),
            )
        )
    return words


def _tables(page: Any) -> list[TableBlock]:
    blocks = []
    for table in page.find_tables():
        rows = [[(cell or "").strip() for cell in row] for row in table.extract()]
        if not rows or max(len(r) for r in rows) < 2 or len(rows) < 2:
            continue  # a single box or line is a frame, not a table
        widest = max(table.rows, key=lambda r: sum(c is not None for c in r.cells))
        widths = [c[2] - c[0] for c in widest.cells if c is not None]
        blocks.append(TableBlock(tuple(table.bbox), rows, widths))
    return blocks


def _pictures(page: Any, tables: list[TableBlock]) -> list[Picture]:
    pictures = []
    width, height = float(page.width), float(page.height)
    for img in page.images:
        box: Box = (
            max(0.0, float(img["x0"])),
            max(0.0, float(img["top"])),
            min(width, float(img["x1"])),
            min(height, float(img["bottom"])),
        )
        if box[2] - box[0] < MIN_PICTURE_PT or box[3] - box[1] < MIN_PICTURE_PT:
            continue
        if any(_contains(t.box, box) for t in tables):
            continue
        pictures.append(Picture(box))
    return pictures


def _looks_scanned(page: Any, words: list[Word]) -> bool:
    """No real text layer: no words, or a page-sized picture with only a few words on it."""
    if not words:
        return True
    area = float(page.width) * float(page.height)
    covered = max(
        (
            (float(i["x1"]) - float(i["x0"])) * (float(i["bottom"]) - float(i["top"]))
            for i in page.images
        ),
        default=0.0,
    )
    return covered >= 0.6 * area and sum(len(w.text) for w in words) < 40


def _contains(outer: Box, inner: Box) -> bool:
    return (
        outer[0] - 1 <= inner[0]
        and outer[1] - 1 <= inner[1]
        and inner[2] <= outer[2] + 1
        and inner[3] <= outer[3] + 1
    )


# ── PDF → Word ────────────────────────────────────────────────────────────────────────────


@contextmanager
def _closing_pdfium(path: Path) -> Iterator[pdfium.PdfDocument]:
    doc = pdfium.PdfDocument(str(path))
    try:
        yield doc
    finally:
        doc.close()


def _set_page(section: Any, width: float, height: float, frame: Box) -> tuple[float, float]:
    """Page size and margins from the content frame; returns the usable (width, height)."""
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    left = max(MIN_MARGIN_PT, min(frame[0], width / 3))
    right = max(MIN_MARGIN_PT, min(width - frame[2], width / 3))
    top = max(MIN_MARGIN_PT, min(frame[1], height / 3))
    bottom = max(MIN_MARGIN_PT, min(height - frame[3], height / 3))
    section.left_margin, section.right_margin = Pt(left), Pt(right)
    section.top_margin, section.bottom_margin = Pt(top), Pt(bottom)
    section.header_distance = section.footer_distance = Pt(MIN_MARGIN_PT / 2)
    return width - left - right, height - top - bottom


def _write_paragraph(doc: Any, para: Paragraph, base_size: float) -> None:
    p = doc.add_paragraph()
    if para.kind == "heading":
        p.style = doc.styles[f"Heading {para.level}"]
    fmt = p.paragraph_format
    fmt.alignment = _ALIGN[para.align]
    fmt.space_before = Pt(min(para.space_before, 72))
    fmt.space_after = Pt(0)
    fmt.keep_with_next = para.kind == "heading"
    if para.left_indent > 1:
        fmt.left_indent = Pt(para.left_indent)
    if abs(para.first_indent) > 1:
        fmt.first_line_indent = Pt(para.first_indent)
    if para.line_pitch and para.line_pitch > para.size:
        fmt.line_spacing = Pt(min(para.line_pitch, 3 * para.size))
    for r in para.runs():
        run = p.add_run(r.text)
        font = run.font
        font.name = r.font
        font.size = Pt(r.size)
        font.bold = r.bold
        font.italic = r.italic
        if r.color:
            font.color.rgb = RGBColor.from_string(r.color)
        elif para.kind == "heading":
            font.color.rgb = RGBColor(0, 0, 0)  # the template's headings are blue


def _write_table(doc: Any, block: TableBlock, space_before: float) -> None:
    if space_before > 2:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(max(1.0, min(space_before, 36)))
    cols = max(len(r) for r in block.rows)
    table = doc.add_table(rows=len(block.rows), cols=cols)
    table.style = doc.styles["Table Grid"]
    for r, row in enumerate(block.rows):
        for c in range(cols):
            table.cell(r, c).text = row[c] if c < len(row) else ""
    if len(block.column_widths) == cols:
        for c, width in enumerate(block.column_widths):
            for cell in table.columns[c].cells:
                cell.width = Pt(width)


def _write_picture(
    doc: Any,
    image: Image.Image,
    picture: Picture,
    usable: tuple[float, float],
    page_width: float,
) -> None:
    x0, top, x1, bottom = picture.box
    scale = IMAGE_DPI / 72
    crop = image.crop((int(x0 * scale), int(top * scale), int(x1 * scale), int(bottom * scale)))
    buffer = io.BytesIO()
    if crop.width * crop.height > 300_000:
        crop.save(buffer, "JPEG", quality=85)
    else:
        crop.save(buffer, "PNG")
    buffer.seek(0)
    width, height = x1 - x0, bottom - top
    fit = min(1.0, (usable[0] - 2) / width, (usable[1] - 4) / height)
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(min(picture.space_before, 72))
    fmt.space_after = Pt(0)
    centre = (x0 + x1) / 2
    if abs(centre - page_width / 2) < 0.05 * page_width:
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif x0 > page_width / 2:
        fmt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(buffer, width=Pt(width * fit), height=Pt(height * fit))


def _convert_to_docx(
    path: Path, shown: Path, dest: Path, ctx: JobContext, ocr_langs: list[str] | None, work: Path
) -> dict[str, int]:
    doc = Document()
    counts = {"pages": 0, "paragraphs": 0, "tables": 0, "images": 0, "ocr_pages": 0}
    plumber = _open_plumber(path, shown)
    try:
        with _closing_pdfium(path) as rendered:
            total = len(plumber.pages)
            all_words: list[list[Word]] = []
            for page in plumber.pages:
                all_words.append(page_words(page))
            base = body_size([w for words in all_words for w in words])
            normal = doc.styles["Normal"].font
            normal.size = Pt(base)
            for index, page in enumerate(plumber.pages):
                ctx.check_cancelled()
                ctx.progress(5 + 85 * index / max(1, total), f"Page {index + 1} of {total}")
                width, height = float(page.width), float(page.height)
                tables = _tables(page)
                pictures = _pictures(page, tables)
                words = [w for w in all_words[index] if not any(inside(w, t.box) for t in tables)]
                image: Image.Image | None = None
                if ocr_langs and _looks_scanned(page, all_words[index]):
                    image = render_page(rendered, index, OCR_DPI)
                    words = ocr_page_words(image, work / "ocr", ocr_langs, ctx, f"p{index}")
                    if words:
                        counts["ocr_pages"] += 1
                        pictures = []  # the scan is replaced by its text
                    image = None
                elements_frame = text_frame(words, width, height)
                boxes = [elements_frame] if words else []
                boxes += [t.box for t in tables] + [p.box for p in pictures]
                frame = (
                    (
                        min(b[0] for b in boxes),
                        min(b[1] for b in boxes),
                        max(b[2] for b in boxes),
                        max(b[3] for b in boxes),
                    )
                    if boxes
                    else (72.0, 72.0, width - 72.0, height - 72.0)
                )
                section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
                usable = _set_page(section, width, height, frame)
                paragraphs = group_paragraphs(group_lines(words), base, frame)
                elements: list[tuple[float, Any]] = [(p.top, p) for p in paragraphs]
                elements += [(t.top, t) for t in tables] + [(p.box[1], p) for p in pictures]
                elements.sort(key=lambda e: e[0])
                previous_bottom = frame[1]
                if pictures:
                    image = render_page(rendered, index, IMAGE_DPI)
                for top, element in elements:
                    gap = max(0.0, top - previous_bottom)
                    if isinstance(element, Paragraph):
                        element.space_before = gap
                        _write_paragraph(doc, element, base)
                        counts["paragraphs"] += 1
                        previous_bottom = element.bottom
                    elif isinstance(element, TableBlock):
                        _write_table(doc, element, gap)
                        counts["tables"] += 1
                        previous_bottom = element.box[3]
                    else:
                        assert image is not None
                        element.space_before = gap
                        _write_picture(doc, image, element, usable, width)
                        counts["images"] += 1
                        previous_bottom = element.box[3]
                counts["pages"] += 1
    finally:
        plumber.close()
    ctx.progress(92, "Saving")
    doc.save(str(dest))
    return counts


def pdf_to_word(params: dict[str, Any], ctx: JobContext) -> ToolResult:
    """PDF → DOCX: text with fonts, sizes, bold/italic, colours and alignment, headings,
    ruled tables and pictures; optional OCR for scanned pages."""
    inputs = pdf_inputs(params)
    options = options_dict(params)
    ocr = option_bool(options, "ocr", False)
    ocr_langs = tesseract.check_languages(_ocr_languages(options)) if ocr else None
    outputs: list[Path] = []
    totals: dict[str, int] = {}
    for n, source in enumerate(inputs):
        ctx.check_cancelled()
        work = ctx.temp_dir / f"doc_{n:03d}"
        work.mkdir(parents=True)
        with readable_copy(source, params, work) as (path, _pages):
            tmp = work / "result" / f"{source.stem[:150]}.docx"
            tmp.parent.mkdir()
            counts = _convert_to_docx(path, source, tmp, ctx, ocr_langs, work)
        for key, value in counts.items():
            totals[key] = totals.get(key, 0) + value
        outputs.append(deliver_converted(tmp, params, source, protected=inputs))
    ctx.progress(100, "Done")
    return ToolResult(outputs, {"files": len(outputs), **totals})


def _ocr_languages(options: dict[str, Any]) -> list[str]:
    raw = options.get("languages", ["eng"])
    if isinstance(raw, str):
        raw = [part for part in raw.replace(",", "+").split("+") if part]
    if not isinstance(raw, list) or not all(isinstance(x, str) for x in raw):
        raise EngineError(ErrorCode.INVALID_INPUT, "Invalid OCR languages")
    return [x.strip() for x in raw][:8]


# ── PDF → Excel ───────────────────────────────────────────────────────────────────────────

_NUMBER = re.compile(r"^[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")
_PERCENT = re.compile(r"^[-+]?\d+(\.\d+)?%$")
_SHEET_INVALID = re.compile(r"[\[\]:*?/\\]")
MAX_SHEETS = 250


def cell_value(text: str) -> tuple[Any, str | None]:
    """Spreadsheet value and number format for a cell's text ('1,234.5' -> 1234.5)."""
    text = text.strip()
    if _NUMBER.match(text):
        plain = text.replace(",", "")
        number: float | int = float(plain) if "." in plain else int(plain)
        return number, "#,##0.##" if "," in text else None
    if _PERCENT.match(text):
        return float(text[:-1]) / 100, "0.##%"
    return text, None


@dataclass
class _SheetRow:
    top: float
    cells: list[tuple[int, str, bool]]  # (column, text, bold)


def _page_rows(page: Any, words: list[Word], tables: list[TableBlock]) -> list[_SheetRow]:
    """A page as spreadsheet rows: ruled tables cell by cell, other text split into columns."""
    loose = [w for w in words if not any(inside(w, t.box) for t in tables)]
    segments = split_segments(group_lines(loose))
    anchors = cluster_columns([s.x0 for s in segments])
    by_line: dict[int, _SheetRow] = {}
    rows: list[_SheetRow] = []
    for seg in segments:
        key = round((seg.top + seg.bottom) / 2)
        row = next((r for k, r in by_line.items() if abs(k - key) <= 3), None)
        if row is None:
            row = _SheetRow(seg.top, [])
            by_line[key] = row
            rows.append(row)
        column = min(range(len(anchors)), key=lambda i: abs(anchors[i] - seg.x0))
        row.cells.append((column, seg.text(), all(w.bold for w in seg.words)))
    for table in tables:
        for r, values in enumerate(table.rows):
            rows.append(
                _SheetRow(table.box[1] + r * 0.01, [(c, v, r == 0) for c, v in enumerate(values)])
            )
    rows.sort(key=lambda r: r.top)
    return rows


def _write_rows(sheet: Any, rows: list[_SheetRow], start: int) -> int:
    """Write rows from 1-based row ``start``; returns the next free row."""
    bold = Font(bold=True)
    r = start
    for row in rows:
        used: set[int] = set()
        for column, text, is_bold in row.cells:
            while column in used:  # two pieces landed in one column: keep both
                column += 1
            used.add(column)
            value, number_format = cell_value(text)
            cell = sheet.cell(row=r, column=column + 1, value=value)
            if number_format:
                cell.number_format = number_format
            if is_bold:
                cell.font = bold
        r += 1
    return r


def _autosize(sheet: Any) -> None:
    widths: dict[str, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                length = len(str(cell.value))
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), length)
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = min(60, max(8, width * 1.15 + 2))


def _sheet_title(name: str, taken: set[str]) -> str:
    base = _SHEET_INVALID.sub(" ", name).strip()[:28] or "Sheet"
    title, n = base, 2
    while title.lower() in taken:
        title = f"{base[:25]} ({n})"
        n += 1
    taken.add(title.lower())
    return title


def _convert_to_xlsx(
    path: Path, shown: Path, dest: Path, ctx: JobContext, layout: str
) -> dict[str, int]:
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    taken: set[str] = set()
    counts = {"pages": 0, "tables": 0, "sheets": 0}
    plumber = _open_plumber(path, shown)
    try:
        pages = list(plumber.pages)
        found: list[tuple[int, list[TableBlock], list[Word]]] = []
        for index, page in enumerate(pages):
            ctx.check_cancelled()
            ctx.progress(5 + 80 * index / max(1, len(pages)), f"Page {index + 1} of {len(pages)}")
            tables = _tables(page)
            counts["tables"] += len(tables)
            found.append((index, tables, page_words(page) if layout != "tables" else []))
            counts["pages"] += 1
        if layout == "tables" and counts["tables"] == 0:
            layout = "pages"  # nothing ruled to detect: export the text in columns instead
            found = [(i, t, page_words(pages[i])) for i, t, _ in found]
        if layout == "tables":
            for index, tables, _ in found:
                for n, table in enumerate(tables, start=1):
                    if len(taken) >= MAX_SHEETS:
                        break
                    sheet = wb.create_sheet(_sheet_title(f"Page {index + 1} Table {n}", taken))
                    _write_rows(sheet, _page_rows(pages[index], [], [table]), 1)
                    _autosize(sheet)
        elif layout == "pages":
            for index, tables, words in found[:MAX_SHEETS]:
                sheet = wb.create_sheet(_sheet_title(f"Page {index + 1}", taken))
                _write_rows(sheet, _page_rows(pages[index], words, tables), 1)
                _autosize(sheet)
        else:
            sheet = wb.create_sheet(_sheet_title(shown.stem or "Sheet", taken))
            r = 1
            for index, tables, words in found:
                r = _write_rows(sheet, _page_rows(pages[index], words, tables), r) + 1
            _autosize(sheet)
    finally:
        plumber.close()
    if not wb.worksheets:
        wb.create_sheet("Sheet1")
    counts["sheets"] = len(wb.worksheets)
    ctx.progress(92, "Saving")
    wb.save(str(dest))
    return counts


def pdf_to_excel(params: dict[str, Any], ctx: JobContext) -> ToolResult:
    """PDF → XLSX: ruled tables (one sheet each), or all content per page / in one sheet."""
    inputs = pdf_inputs(params)
    layout = option_choice(options_dict(params), "layout", ("tables", "pages", "single"), "tables")
    outputs: list[Path] = []
    totals: dict[str, int] = {}
    for n, source in enumerate(inputs):
        ctx.check_cancelled()
        work = ctx.temp_dir / f"doc_{n:03d}"
        work.mkdir(parents=True)
        with readable_copy(source, params, work) as (path, _pages):
            tmp = work / "result" / f"{source.stem[:150]}.xlsx"
            tmp.parent.mkdir()
            counts = _convert_to_xlsx(path, source, tmp, ctx, layout)
        for key, value in counts.items():
            totals[key] = totals.get(key, 0) + value
        outputs.append(deliver_converted(tmp, params, source, protected=inputs))
    ctx.progress(100, "Done")
    return ToolResult(outputs, {"files": len(outputs), **totals})


# ── PDF → PowerPoint ──────────────────────────────────────────────────────────────────────

EMU_PER_PT = 12700
SLIDE_MIN_PT, SLIDE_MAX_PT = 72.0, 4032.0  # PowerPoint accepts 1..56 inches
SLIDE_DPI = 150


def _strip_text_operators(pdf: pikepdf.Pdf) -> None:
    """Remove all text-showing blocks (BT ... ET) from pages and the forms they use, so the
    rendered background carries pictures and drawings but no text."""
    seen: set[tuple[int, int]] = set()

    def strip(owner: Any) -> None:
        kept = []
        depth = 0
        for instruction in pikepdf.parse_content_stream(owner):
            name = str(instruction.operator)
            if name == "BT":
                depth += 1
                continue
            if name == "ET":
                depth = max(0, depth - 1)
                continue
            if depth == 0:
                kept.append(instruction)
        data = pikepdf.unparse_content_stream(kept)
        if isinstance(owner, pikepdf.Page):
            owner.obj.Contents = pdf.make_stream(data)
        else:
            owner.write(data)

    def visit_resources(resources: Any) -> None:
        xobjects = resources.get("/XObject") if resources is not None else None
        if xobjects is None:
            return
        for _name, xobj in xobjects.items():
            if xobj.get("/Subtype") != "/Form" or not xobj.is_indirect:
                continue
            if xobj.objgen in seen:
                continue
            seen.add(xobj.objgen)
            strip(xobj)
            visit_resources(xobj.get("/Resources"))

    for page in pdf.pages:
        page.contents_coalesce()
        strip(page)
        visit_resources(page.obj.get("/Resources"))


def _slide_size(width: float, height: float) -> tuple[float, float]:
    scale = min(1.0, SLIDE_MAX_PT / max(width, height))
    scale = max(scale, SLIDE_MIN_PT / max(1.0, min(width, height)))
    return width * scale, height * scale


def _emu(points: float) -> Emu:
    return Emu(int(points * EMU_PER_PT))


def _add_text_box(slide: Any, para: Paragraph, scale: float, offset: tuple[float, float]) -> None:
    x0 = min(line.x0 for line in para.lines)
    x1 = max(line.x1 for line in para.lines)
    # A little slack: PowerPoint's text metrics differ slightly from the PDF's.
    width = (x1 - x0) * 1.06 + 2
    height = para.bottom - para.top
    box = slide.shapes.add_textbox(
        _emu(offset[0] + x0 * scale),
        _emu(offset[1] + para.top * scale),
        _emu(width * scale),
        _emu(height * scale),
    )
    frame = box.text_frame
    frame.word_wrap = len(para.lines) > 1
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = {
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get(para.align, PP_ALIGN.LEFT)
    if para.line_pitch:
        p.line_spacing = PptPt(para.line_pitch * scale)
    for r in para.runs():
        run = p.add_run()
        run.text = r.text.replace("\t", " ")
        font = run.font
        font.name = r.font
        font.size = PptPt(max(1.0, r.size * scale))
        font.bold = r.bold
        font.italic = r.italic
        if r.color:
            font.color.rgb = PptRGB.from_string(r.color)  # type: ignore[no-untyped-call]


def _convert_to_pptx(
    path: Path, shown: Path, dest: Path, ctx: JobContext, work: Path
) -> dict[str, int]:
    textless = work / "textless.pdf"
    try:
        with pikepdf.open(path) as pdf:
            _strip_text_operators(pdf)
            pdf.save(textless)
    except pikepdf.PdfError as exc:
        raise EngineError(
            ErrorCode.CORRUPT_FILE, f"{shown.name} is damaged and could not be read", str(shown)
        ) from exc
    deck = Presentation()
    blank = deck.slide_layouts[6]
    counts = {"pages": 0, "text_boxes": 0}
    plumber = _open_plumber(path, shown)
    try:
        with _closing_pdfium(textless) as rendered:
            total = len(plumber.pages)
            sizes = [(float(p.width), float(p.height)) for p in plumber.pages]
            slide_w, slide_h = _slide_size(*sizes[0])
            deck.slide_width = _emu(slide_w)
            deck.slide_height = _emu(slide_h)
            for index, page in enumerate(plumber.pages):
                ctx.check_cancelled()
                ctx.progress(5 + 85 * index / max(1, total), f"Page {index + 1} of {total}")
                width, height = sizes[index]
                # Pages of another size are fitted into the deck's slide size, centred.
                scale = min(slide_w / width, slide_h / height)
                offset = ((slide_w - width * scale) / 2, (slide_h - height * scale) / 2)
                slide = deck.slides.add_slide(blank)
                background = render_page(rendered, index, SLIDE_DPI)
                buffer = io.BytesIO()
                background.save(buffer, "PNG", optimize=True)
                buffer.seek(0)
                slide.shapes.add_picture(
                    buffer,
                    _emu(offset[0]),
                    _emu(offset[1]),
                    _emu(width * scale),
                    _emu(height * scale),
                )
                words = page_words(page)
                lines = split_segments(group_lines(words), factor=2.0)
                frame = text_frame(words, width, height)
                for para in group_paragraphs(lines, body_size(words), frame):
                    _add_text_box(slide, para, scale, offset)
                    counts["text_boxes"] += 1
                counts["pages"] += 1
    finally:
        plumber.close()
    ctx.progress(92, "Saving")
    deck.save(str(dest))
    return counts


def pdf_to_powerpoint(params: dict[str, Any], ctx: JobContext) -> ToolResult:
    """PDF → PPTX: one slide per page — the page's pictures and drawings as the background,
    its text as editable text boxes on top."""
    inputs = pdf_inputs(params)
    outputs: list[Path] = []
    totals: dict[str, int] = {}
    for n, source in enumerate(inputs):
        ctx.check_cancelled()
        work = ctx.temp_dir / f"doc_{n:03d}"
        work.mkdir(parents=True)
        with readable_copy(source, params, work) as (path, _pages):
            tmp = work / "result" / f"{source.stem[:150]}.pptx"
            tmp.parent.mkdir()
            counts = _convert_to_pptx(path, source, tmp, ctx, work)
        for key, value in counts.items():
            totals[key] = totals.get(key, 0) + value
        outputs.append(deliver_converted(tmp, params, source, protected=inputs))
    ctx.progress(100, "Done")
    return ToolResult(outputs, {"files": len(outputs), **totals})


# ── PDF → Text ────────────────────────────────────────────────────────────────────────────


def _page_text(page: Any, keep_layout: bool) -> str:
    if keep_layout:
        text = page.extract_text(layout=True) or ""
        return "\n".join(line.rstrip() for line in text.splitlines()).strip("\n")
    return "\n".join(line.text() for line in group_lines(page_words(page)))


def pdf_to_text(params: dict[str, Any], ctx: JobContext) -> ToolResult:
    """PDF → UTF-8 .txt (Windows line endings); pages separated by a form feed."""
    inputs = pdf_inputs(params)
    keep_layout = option_bool(options_dict(params), "keep_layout", False)
    outputs: list[Path] = []
    pages_total = 0
    chars = 0
    for n, source in enumerate(inputs):
        ctx.check_cancelled()
        work = ctx.temp_dir / f"doc_{n:03d}"
        work.mkdir(parents=True)
        texts: list[str] = []
        with readable_copy(source, params, work) as (path, _pages):
            plumber = _open_plumber(path, source)
            try:
                total = len(plumber.pages)
                for index, page in enumerate(plumber.pages):
                    ctx.check_cancelled()
                    done = (n + index / max(1, total)) / len(inputs)
                    ctx.progress(95 * done, f"Page {index + 1} of {total}")
                    texts.append(_page_text(page, keep_layout))
            finally:
                plumber.close()
        tmp = work / "result" / f"{source.stem[:150]}.txt"
        tmp.parent.mkdir()
        tmp.write_text("\n\f".join(texts) + "\n", encoding="utf-8", newline="\r\n")
        pages_total += len(texts)
        chars += sum(len(t) for t in texts)
        outputs.append(deliver_converted(tmp, params, source, protected=inputs))
    ctx.progress(100, "Done")
    return ToolResult(outputs, {"files": len(outputs), "pages": pages_total, "characters": chars})
